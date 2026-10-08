"""提取已保存的搜索结果，生成详情网址，保存并核对后删除源 MHTML。"""

import base64
import argparse
import hashlib
import os
from pathlib import Path
import tempfile

from lxml import html
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from medicine_nmpa import read_page_source, check_cleanup_target


DOMESTIC_ITEM_ID = "ff80808183cad75001840881f848179f"
OVERSEAS_ITEM_ID = "ff80808183cad7500184088665711800"
ROOT = Path(__file__).resolve().parent / "medicine_nmpa"
INPUT_NAMES = ("国家药品监督管理数据查询.mhtml", "国家药品监督管理局数据查询.mhtml")
OUTPUT_NAME = "查询结果整理.xlsx"
DETAIL_URL = "https://datasearch.nmpa.gov.cn/datasearch/search-info.html"


def make_detail_url(approval_number, item_id=DOMESTIC_ITEM_ID):
    record_id = hashlib.md5(approval_number.encode("utf-8")).hexdigest()
    parameters = f"id={record_id}&itemId={item_id}"
    encoded = base64.b64encode(parameters.encode("utf-8")).decode("ascii")
    return f"{DETAIL_URL}?nmpa={encoded}"


def extract_results(path):
    source, location = read_page_source(path)
    document = html.fromstring(source)
    headers = document.xpath('//table[contains(concat(" ", normalize-space(@class), " "), " el-table__header ")]')
    bodies = document.xpath('//table[contains(concat(" ", normalize-space(@class), " "), " el-table__body ")]')
    if len(headers) != 1 or len(bodies) != 1:
        raise ValueError("未找到唯一的搜索结果表，请确认保存的是已加载完成的查询结果页。")
    columns = [" ".join(cell.text_content().split()) for cell in headers[0].xpath(".//th") if cell.text_content().strip()]
    number_columns = [name for name in ("批准文号", "注册证号") if name in columns]
    if len(number_columns) != 1 or "详情" not in columns or len(set(columns)) != len(columns):
        raise ValueError("搜索结果表头需包含唯一的批准文号/注册证号列及详情列，且列名不能重复，请复查。")
    number_column = number_columns[0]
    # 根据原始表头选择数据库，不能只用编号前缀判断境内或境外。
    item_id = OVERSEAS_ITEM_ID if number_column == "注册证号" else DOMESTIC_ITEM_ID
    number_index, link_index = columns.index(number_column), columns.index("详情")
    columns[link_index] = "详情链接"
    records, seen = [], set()
    for row in bodies[0].xpath(".//tbody/tr"):
        cells = row.xpath("./td")
        if not cells:
            continue
        if len(cells) != len(columns):
            raise ValueError("搜索结果行的列数与表头不一致，请复查。")
        values = []
        for cell in cells:
            value = " ".join(cell.text_content().split())
            # 使用已保存的悬浮提示补全页面上被省略的名称。
            for element in cell.xpath(".//*[@aria-describedby]"):
                for identifier in element.get("aria-describedby").split():
                    tips = document.xpath("//*[@id=$identifier]", identifier=identifier)
                    if len(tips) == 1 and tips[0].text_content().strip():
                        value = " ".join(tips[0].text_content().split())
            values.append(value)
        number = values[number_index]
        if not number or number in seen:
            raise ValueError(f"{number_column}为空或重复，请复查：{number or '空值'}")
        seen.add(number)
        values[link_index] = make_detail_url(number, item_id)
        records.append(values)
    if not records:
        raise ValueError("搜索结果表没有数据，未修改 Excel，也未删除源文件。")
    return columns, records, location


def save_results(output, columns, records, location):
    # 详情链接紧跟批准文号或注册证号，便于逐条打开网页。
    ordered_columns = [column for column in columns if column != "详情链接"]
    number_column = "注册证号" if "注册证号" in columns else "批准文号"
    ordered_columns.insert(ordered_columns.index(number_column) + 1, "详情链接")
    indices = [columns.index(column) for column in ordered_columns]
    records = [[record[index] for index in indices] for record in records]
    columns = ordered_columns
    book = Workbook()
    sheet = book.active
    sheet.title = "查询结果"
    sheet.append(columns)
    for record in records:
        sheet.append(record)
    link_index = columns.index("详情链接") + 1
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[link_index - 1].hyperlink = row[link_index - 1].value
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="245A81")
    for index, name in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 50 if name in {"详情链接", "生产单位"} else 24
    sheet["A1"].comment = Comment(f"来源：{location}\n仅包含此文件中保存的当前页结果，不包含其他分页。", "提取说明")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        checked = load_workbook(temporary)
        try:
            expected = [tuple(columns)] + [tuple(value or None for value in row) for row in records]
            if list(checked["查询结果"].values) != expected:
                raise ValueError("导出的搜索结果核对失败，源文件保留。")
        finally:
            checked.close()
        os.replace(temporary, output)
    finally:
        book.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()


def run(directory=ROOT):
    root = Path(directory).resolve(strict=True)
    inputs = [root / name for name in INPUT_NAMES if (root / name).exists()]
    if not inputs:
        print("未找到搜索结果 MHTML，未修改任何文件。")
        return
    if len(inputs) != 1:
        raise ValueError("发现两个候选搜索结果文件，请只保留本次需要整理的一个。")
    source = inputs[0]
    check_cleanup_target(source, root)
    columns, records, location = extract_results(source)
    output = root / OUTPUT_NAME
    save_results(output, columns, records, location)
    print(f"已保存 {len(records)} 条当前页搜索结果及详情链接：{output}")
    try:
        check_cleanup_target(source, root)
        source.unlink()
    except (OSError, ValueError):
        raise ValueError(f"Excel 已保存，但源文件未能删除，请手动处理：{source}") from None
    print(f"已删除搜索结果源文件：{source.name}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, add_help=False, usage="python medicine_nmpa_search.py")
    parser.add_argument("-h", "--help", action="help", help="显示帮助并退出")
    args = parser.parse_args()
    try:
        run()
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else f"操作失败（{type(exc).__name__}），请检查网页格式、文件权限及 Excel 是否已关闭。"
        print("已停止：" + message)
        raise SystemExit(1)
