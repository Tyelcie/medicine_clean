"""连接手动启动的 Chrome，按查询结果逐页保存 MHTML，每页完成后休息 5 秒。"""

import os
from pathlib import Path
import re
import tempfile
import time
from urllib.parse import urlparse

from openpyxl import load_workbook
from playwright.sync_api import Error, TimeoutError, sync_playwright

from medicine_nmpa import extract_page


ROOT = Path(__file__).resolve().parent / "medicine_nmpa"
CDP_URL = "http://127.0.0.1:9222"
PAGE_INTERVAL_SECONDS = 5
MANUAL_VERIFY_TIMEOUT_SECONDS = 600


def read_tasks(path):
    book = load_workbook(path, read_only=True)
    try:
        if "查询结果" not in book.sheetnames:
            raise ValueError("Excel 中没有“查询结果”工作表。")
        rows = iter(book["查询结果"].values)
        columns = list(next(rows, ()))
        if "批准文号" not in columns or "详情链接" not in columns:
            raise ValueError("查询结果表缺少批准文号或详情链接列。")
        tasks, seen = [], set()
        for values in rows:
            if all(value is None for value in values):
                continue
            record = dict(zip(columns, values))
            number = str(record.get("批准文号") or "").strip()
            url = str(record.get("详情链接") or "").strip()
            if not re.fullmatch(r"国药准字[A-Za-z0-9]+", number):
                raise ValueError(f"批准文号格式不正确，请复查：{number or '空值'}")
            parsed = urlparse(url)
            if parsed.scheme != "https" or parsed.netloc != "datasearch.nmpa.gov.cn" or parsed.path != "/datasearch/search-info.html":
                raise ValueError(f"{number} 的详情链接为空或不属于药监局详情页。")
            if number in seen:
                raise ValueError(f"查询结果存在重复批准文号，请复查：{number}")
            seen.add(number)
            tasks.append((number, url))
        return tasks
    finally:
        book.close()


def save_snapshot(page, number, target):
    session = page.context.new_cdp_session(page)
    try:
        snapshot = session.send("Page.captureSnapshot", {"format": "mhtml"})["data"]
    finally:
        session.detach()
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".mhtml", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(snapshot.encode("utf-8"))
        record = extract_page(temporary)
        if record["批准文号"] != number:
            raise ValueError(f"页面批准文号与目标不一致，未保存：{number}")
        if target.exists():
            raise ValueError(f"目标文件已经存在，未覆盖：{target.name}")
        os.rename(temporary, target)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def run(directory=ROOT):
    root = Path(directory).resolve(strict=True)
    tasks = read_tasks(root / "查询结果整理.xlsx")
    pending = []
    for number, url in tasks:
        target = root / f"{number}.mhtml"
        if target.exists():
            if target.is_symlink() or extract_page(target)["批准文号"] != number:
                raise ValueError(f"已有文件异常，请复查：{target.name}")
            print(f"已有有效网页，跳过：{number}")
        else:
            pending.append((number, url, target))
    if not pending:
        print("没有需要保存的新网页。")
        return
    with sync_playwright() as playwright:
        try:
            browser = playwright.chromium.connect_over_cdp(CDP_URL, no_defaults=True, timeout=10000)
        except Error:
            raise ValueError("无法连接 Chrome。请先运行启动命令，开启本机 9222 调试端口，再运行本脚本。") from None
        if not browser.contexts:
            raise ValueError("Chrome 没有可用的浏览器会话。")
        # 只新建一个下载标签页，不操作其他标签页，也不关闭用户的浏览器。
        page = browser.contexts[0].new_page()
        for index, (number, url, target) in enumerate(pending, 1):
            print(f"正在打开第 {index}/{len(pending)} 页：{number}", flush=True)
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=45000)
            except TimeoutError:
                print("导航等待超时，继续等待药品详情显示。", flush=True)
            page.bring_to_front()
            print("如有滑动验证，请手动完成；详情加载后自动保存，最多等待 10 分钟。", flush=True)
            try:
                page.wait_for_function("""number => Array.from(document.querySelectorAll('table')).some(table => {
                    const fields = Object.fromEntries(Array.from(table.querySelectorAll('tr'))
                        .map(row => Array.from(row.cells, cell => cell.innerText.trim()))
                        .filter(cells => cells.length === 2));
                    return fields['批准文号'] === number && Boolean(fields['产品名称']);
                })""", arg=number, timeout=MANUAL_VERIFY_TIMEOUT_SECONDS * 1000, polling=1000)
            except TimeoutError:
                raise ValueError(f"{number} 等待详情超时，已停止；已保存文件保留，可稍后重新运行续传。") from None
            save_snapshot(page, number, target)
            print(f"已保存并核对：{target.name}；休息 {PAGE_INTERVAL_SECONDS} 秒。", flush=True)
            time.sleep(PAGE_INTERVAL_SECONDS)
        print(f"完成，本次保存 {len(pending)} 个 MHTML 文件。")


if __name__ == "__main__":
    try:
        run()
    except KeyboardInterrupt:
        print("已手动停止，已保存网页保留。")
        raise SystemExit(1)
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else f"保存失败（{type(exc).__name__}），请检查浏览器连接、Excel 和目录权限；已完成文件保留。"
        print("已停止：" + message)
        raise SystemExit(1)
