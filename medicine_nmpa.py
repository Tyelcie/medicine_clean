"""将已保存的药监局详情网页追加到 Excel，核对成功后清理源网页。"""

import argparse
from collections import Counter
from email import policy
from email.parser import BytesParser
import os
from pathlib import Path
import re
import shutil
import tempfile

from lxml import html
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent / "medicine_nmpa"
REMOVED_COLUMNS = {"来源文件", "整理说明"}
WEB_SUFFIXES = {".html", ".htm", ".mhtml", ".mht"}
OVERSEAS_FIELD_ALIASES = {
    "批准文号": ("注册证号",),
    "产品名称": ("产品名称（中文）", "产品名称（英文）"),
    "英文名称": ("产品名称（英文）",),
    "商品名": ("商品名（中文）", "商品名（英文）"),
    "剂型": ("剂型（中文）",),
    "规格": ("规格（中文）",),
    "上市许可持有人": ("上市许可持有人中文名称", "上市许可持有人英文名称"),
    "上市许可持有人地址": ("上市许可持有人地址（中文）", "上市许可持有人地址（英文）"),
    "生产单位": ("生产厂商（中文）", "生产厂商（英文）"),
    "生产地址": ("厂商地址（中文）", "厂商地址（英文）"),
}


def duplicate_numbers(numbers):
    return sorted(number for number, count in Counter(numbers).items() if count > 1)


def read_existing(path):
    if not path.exists():
        return [], []
    with path.open("rb") as stream:
        book = load_workbook(stream)
        try:
            if "药品信息" not in book.sheetnames:
                raise ValueError("Excel 中缺少“药品信息”工作表，请复查。")
            rows = list(book["药品信息"].values)
        finally:
            book.close()
    if not rows or "批准文号" not in rows[0]:
        raise ValueError("表头缺少“批准文号”，请复查。")
    columns = list(rows[0])
    if any(not column for column in columns) or len(set(columns)) != len(columns):
        raise ValueError("Excel 表头存在空列名或重复列名，请复查。")
    records = []
    for line, values in enumerate(rows[1:], 2):
        if all(value is None for value in values):
            continue
        record = dict(zip(columns, values))
        number = str(record.get("批准文号") or "").strip()
        if not number or number.startswith("="):
            raise ValueError(f"Excel 第 {line} 行批准文号为空或为公式，请复查。")
        record["批准文号"] = number
        records.append(record)
    duplicates = duplicate_numbers(record["批准文号"] for record in records)
    if duplicates:
        raise ValueError("Excel 已有重复批准文号，请复查：" + "、".join(duplicates))
    return [column for column in columns if column not in REMOVED_COLUMNS], records


def read_page_source(path):
    if path.suffix.lower() not in {".mhtml", ".mht"}:
        source = path.read_text(encoding="utf-8-sig")
        match = re.search(r"<!-- saved from url=\(\d+\)(.*?) -->", source)
        return source, match.group(1) if match else ""
    # 解码单文件网页的 MIME 内容，不执行其中的脚本或加载外部资源。
    archive = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    parts = [part for part in archive.walk() if part.get_content_type() == "text/html"]
    location = str(archive.get("Snapshot-Content-Location", "")).strip()
    start = archive.get_param("start")
    selected = [part for part in parts if start and part.get("Content-ID") == start]
    if not selected:
        selected = [part for part in parts if location and part.get("Content-Location") == location]
    if not selected and len(parts) == 1:
        selected = parts
    if len(selected) != 1:
        raise ValueError(f"{path.name} 未找到唯一的主 HTML 页面，请复查保存内容。")
    part = selected[0]
    payload = part.get_payload(decode=True)
    if not payload:
        raise ValueError(f"{path.name} 的主页面内容为空，请重新保存。")
    document = html.fromstring(payload, parser=html.HTMLParser(encoding=part.get_content_charset()))
    return html.tostring(document, encoding="unicode"), str(part.get("Content-Location") or location)


def extract_page(path):
    source, location = read_page_source(path)
    document = html.fromstring(source)
    candidates = []
    for table in document.xpath("//table"):
        fields = {}
        for row in table.xpath(".//tr"):
            cells = row.xpath("./th|./td")
            if len(cells) != 2:
                continue
            key, value = [" ".join(cell.text_content().split()) for cell in cells]
            if not key:
                continue
            if key in fields:
                raise ValueError(f"{path.name} 的表格字段重复：{key}，请复查。")
            fields[key] = value
        domestic = fields.get("批准文号") and fields.get("产品名称")
        overseas = fields.get("注册证号") and (fields.get("产品名称（中文）") or fields.get("产品名称（英文）"))
        if domestic or overseas:
            candidates.append(fields)
    if len(candidates) != 1:
        raise ValueError(f"{path.name} 未找到唯一的药品详情表，请在详情加载完成后保存“网页，全部”或 MHTML。")
    record = candidates[0]
    if record.get("注册证号"):
        # 保留境外页面所有原始字段，只补充用于统一整理和查重的公共列。
        # 中文缺失或仅为横线占位时使用英文；发证日期不混作批准日期。
        for common_name, original_names in OVERSEAS_FIELD_ALIASES.items():
            if record.get(common_name):
                continue
            record[common_name] = next(
                (record[name] for name in original_names
                 if record.get(name) and not re.fullmatch(r"[-—–－\s]+", record[name])),
                "",
            )
    if not record.get("批准文号") or not record.get("产品名称"):
        raise ValueError(f"{path.name} 缺少有效的批准文号/注册证号或产品名称，请复查。")
    match = re.fullmatch(r"国药准字[A-Za-z0-9]+", path.stem)
    if match and path.stem != record["批准文号"]:
        raise ValueError(f"{path.name} 的文件名与网页批准文号不一致，请复查。")
    record["来源网址"] = location
    return record


def check_cleanup_target(path, root):
    # 只清理当前目录的源网页及同名附件，拒绝通过链接指向其他位置。
    if path.is_symlink() or path.is_junction() or path.resolve().parent != root:
        raise ValueError(f"清理目标不在指定目录内或为链接，已停止：{path}")
    if path.is_dir():
        for folder, dirs, files in os.walk(path, followlinks=False):
            for name in dirs + files:
                child = Path(folder) / name
                if child.is_symlink() or child.is_junction():
                    raise ValueError(f"附件目录中存在链接，已停止：{child}")


def write_book(path, columns, records):
    # 先写临时文件并重新读取核对，最后原子替换正式文件。
    book = load_workbook(path) if path.exists() else Workbook()
    if not path.exists():
        book.remove(book.active)
    if "药品信息" in book.sheetnames:
        book.remove(book["药品信息"])
    if "文件检查" in book.sheetnames:
        book.remove(book["文件检查"])
    sheet = book.create_sheet("药品信息", 0)
    sheet.append(columns)
    for record in records:
        sheet.append([record.get(column) for column in columns])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="245A81")
    for i, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 45 if "地址" in column or "网址" in column else 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        checked = load_workbook(temporary)
        try:
            actual = list(checked["药品信息"].values)
            expected = [tuple(columns)] + [tuple(record.get(column) if record.get(column) != "" else None for column in columns) for record in records]
            if actual != expected:
                raise ValueError("Excel 写入后的内容核对失败，源网页未删除。")
        finally:
            checked.close()
        os.replace(temporary, path)
    finally:
        book.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()


def run(directory=ROOT, check_only=False):
    root = Path(directory).resolve(strict=True)
    output = root / "药品信息整理.xlsx"
    columns, existing = read_existing(output)
    pages = sorted(path for path in root.iterdir() if path.suffix.lower() in WEB_SUFFIXES and path.is_file())
    if not pages:
        print("没有待处理的 HTML 或 MHTML 网页，未修改 Excel。")
        return
    incoming, targets = [], []
    for page in pages:
        check_cleanup_target(page, root)
        for suffix in (("_files", ".files") if page.suffix.lower() in {".html", ".htm"} else ()):
            folder = page.with_name(page.stem + suffix)
            if folder.exists():
                check_cleanup_target(folder, root)
                if not folder.is_dir():
                    raise ValueError(f"附件路径不是目录，请复查：{folder}")
                targets.append(folder)
        targets.append(page)
        incoming.append(extract_page(page))
    duplicates = duplicate_numbers([record["批准文号"] for record in existing + incoming])
    if duplicates:
        locations = []
        for number in duplicates:
            sources = [page.name for page, record in zip(pages, incoming) if record["批准文号"] == number]
            if any(record["批准文号"] == number for record in existing):
                sources.insert(0, "已有 Excel")
            locations.append(f"{number}：{'、'.join(sources)}")
        raise ValueError("批准文号重复，未写入或删除任何文件，请复查：\n" + "\n".join(locations))
    for record in incoming:
        for column in record:
            if column not in columns and column not in REMOVED_COLUMNS:
                columns.append(column)
    if check_only:
        print(f"检查通过：可追加 {len(incoming)} 条；本次不写入 Excel、不删除网页。")
        return
    write_book(output, columns, existing + incoming)
    print(f"已追加 {len(incoming)} 条药品信息，累计 {len(existing) + len(incoming)} 条：{output}")
    for target in dict.fromkeys(targets):
        try:
            check_cleanup_target(target, root)
            if target.is_dir():
                shutil.rmtree(target)
            else:
                target.unlink()
            print(f"已删除：{target.name}")
        except (OSError, ValueError):
            raise ValueError(f"Excel 已保存，但清理失败：{target}。请手动复查并清理；再次运行会因重复批准文号停止。") from None


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, add_help=False, usage="python medicine_nmpa.py [--check] [--directory 目录]")
    parser.add_argument("-h", "--help", action="help", help="显示帮助并退出")
    parser.add_argument("--check", action="store_true", help="只检查，不写入和删除文件")
    parser.add_argument("--directory", type=Path, default=ROOT, help="存放网页和汇总表的目录")
    args = parser.parse_args()
    try:
        run(args.directory, args.check)
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else f"操作失败（{type(exc).__name__}），请检查文件格式、权限及 Excel 是否已关闭。"
        print("已停止：" + message)
        raise SystemExit(1)
